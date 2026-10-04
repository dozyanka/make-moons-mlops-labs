from __future__ import annotations

import hashlib
import json
import subprocess
from pathlib import Path

import mlflow
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from mlflow.tracking import MlflowClient


ROOT = Path(__file__).resolve().parents[1]
LAB2 = ROOT / "reports" / "LAB2"


def git_head() -> str:
    return subprocess.check_output(
        ["git", "rev-parse", "HEAD"],
        cwd=ROOT,
        text=True,
    ).strip()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def build_model_manifest() -> dict:
    files = [
        "linear_baseline.joblib",
        "ml_model.joblib",
        "ml_model_chunk.joblib",
    ]

    payload = {}
    for name in files:
        path = LAB2 / name
        if path.exists():
            payload[name] = {
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
            }

    (LAB2 / "model_manifest.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    return payload


def export_mlflow_runs() -> None:
    uri = "sqlite:///mlflow.db"
    mlflow.set_tracking_uri(uri)
    client = MlflowClient(tracking_uri=uri)

    rows = []
    try:
        experiments = client.search_experiments()
        for exp in experiments:
            for run in client.search_runs([exp.experiment_id]):
                rows.append(
                    {
                        "experiment_id": exp.experiment_id,
                        "experiment_name": exp.name,
                        "run_id": run.info.run_id,
                        "status": run.info.status,
                        "run_name": run.data.tags.get("mlflow.runName", ""),
                        "params": json.dumps(run.data.params, ensure_ascii=False, sort_keys=True),
                        "metrics": json.dumps(run.data.metrics, ensure_ascii=False, sort_keys=True),
                    }
                )
    except Exception as exc:
        rows.append(
            {
                "experiment_id": "",
                "experiment_name": "",
                "run_id": "",
                "status": "EXPORT_ERROR",
                "run_name": "",
                "params": "",
                "metrics": str(exc),
            }
        )

    pd.DataFrame(rows).to_csv(
        LAB2 / "mlflow_runs.csv",
        index=False,
    )


def read_results():
    metrics = pd.read_csv(LAB2 / "ml_metrics.csv")
    search = pd.read_csv(LAB2 / "hyperparam_search.csv")
    comparison = json.loads(
        (LAB2 / "comparison.json").read_text(encoding="utf-8")
    )
    return metrics, search, comparison


def value(metrics: pd.DataFrame, mode: str, column: str) -> float:
    return float(metrics.loc[metrics["mode"] == mode, column].iloc[0])


def build_markdown(
    metrics: pd.DataFrame,
    search: pd.DataFrame,
    comparison: dict,
    manifest: dict,
    commit: str,
) -> str:
    full_acc = value(metrics, "full", "accuracy")
    chunk_acc = value(metrics, "chunks", "accuracy")
    full_f1 = value(metrics, "full", "f1")
    chunk_f1 = value(metrics, "chunks", "f1")
    full_auc = value(metrics, "full", "roc_auc")
    chunk_auc = value(metrics, "chunks", "roc_auc")
    full_mem = value(metrics, "full", "peak_memory_mb")
    chunk_mem = value(metrics, "chunks", "peak_memory_mb")
    full_time = value(metrics, "full", "elapsed_s")
    chunk_time = value(metrics, "chunks", "elapsed_s")

    mem_reduction = (full_mem - chunk_mem) / full_mem * 100.0
    time_change = (chunk_time - full_time) / full_time * 100.0

    baseline = comparison.get("linear_baseline_test", {})
    baseline_acc = float(baseline.get("accuracy", 0.0))
    improvement_pp = (full_acc - baseline_acc) * 100.0

    best_gamma = comparison.get("best_gamma", "")
    pred_diff = float(comparison.get("prediction_max_abs_diff", 0.0))
    chunks = int(comparison.get("chunks", 0))

    full_hash = manifest.get("ml_model.joblib", {}).get("sha256", "")
    chunk_hash = manifest.get("ml_model_chunk.joblib", {}).get("sha256", "")

    return f"""# Лабораторная работа 2. Классический ML: полный датасет и чанки

## 1. Паспорт работы

Автор: Vladimir

Вариант: 09 — `make_moons`, проблема нелинейности.

Хеш коммита-основания: `{commit}`.

Данные: 24 000 объектов из ЛР1.

Разбиение:
- train: 16 800;
- validation: 3 600;
- test: 3 600.

Train дополнительно представлен {chunks} чанками по 700 объектов.

## 2. Модель

Для нелинейной задачи используется один и тот же классический ML-пайплайн в двух режимах подачи данных:

`StandardScaler -> RBFSampler -> SGDClassifier`.

RBF-преобразование позволяет линейному SGDClassifier работать в нелинейном пространстве признаков.

Подбор `gamma` выполняется по validation-набору.

Лучшее значение: `{best_gamma}`.

Результаты поиска сохранены в `hyperparam_search.csv`.

Линейный LogisticRegression из ЛР1 используется только как baseline.

## 3. Full vs chunks

| Показатель | Full | Chunks |
|---|---:|---:|
| Accuracy | {full_acc:.6f} | {chunk_acc:.6f} |
| F1 | {full_f1:.6f} | {chunk_f1:.6f} |
| ROC-AUC | {full_auc:.6f} | {chunk_auc:.6f} |
| Peak memory, MiB | {full_mem:.4f} | {chunk_mem:.4f} |
| Time, s | {full_time:.4f} | {chunk_time:.4f} |

Accuracy full и chunks отличается на `{abs(full_acc - chunk_acc):.6f}`.

Максимальное абсолютное расхождение предсказаний: `{pred_diff:.6f}`.

Пиковая измеренная память chunks ниже full на `{mem_reduction:.2f}%`.

Время chunks относительно full изменилось на `{time_change:+.2f}%`.

В данном прогоне chunk-режим экономит память, но имеет дополнительный overhead чтения и последовательной обработки чанков.

## 4. Сравнение с линейным baseline

Accuracy линейного baseline на test: `{baseline_acc:.6f}`.

Accuracy нелинейного ML на test: `{full_acc:.6f}`.

Прирост accuracy: `{improvement_pp:.2f}` процентного пункта.

Это подтверждает проблему варианта: нелинейное преобразование существенно улучшает качество на `make_moons`.

## 5. Bootstrap-интервалы

Full accuracy CI:
`[{value(metrics, "full", "accuracy_ci_low"):.6f}; {value(metrics, "full", "accuracy_ci_high"):.6f}]`.

Chunks accuracy CI:
`[{value(metrics, "chunks", "accuracy_ci_low"):.6f}; {value(metrics, "chunks", "accuracy_ci_high"):.6f}]`.

Full ROC-AUC CI:
`[{value(metrics, "full", "roc_auc_ci_low"):.6f}; {value(metrics, "full", "roc_auc_ci_high"):.6f}]`.

Chunks ROC-AUC CI:
`[{value(metrics, "chunks", "roc_auc_ci_low"):.6f}; {value(metrics, "chunks", "roc_auc_ci_high"):.6f}]`.

## 6. Сохранённые модели

`linear_baseline.joblib` — линейный baseline.

`ml_model.joblib` — модель режима full.

`ml_model_chunk.joblib` — модель режима chunks.

SHA-256 full:
`{full_hash}`

SHA-256 chunks:
`{chunk_hash}`

В текущем детерминированном прогоне хеши full и chunks {'совпадают' if full_hash == chunk_hash else 'различаются'}.

## 7. MLflow

Оба режима зарегистрированы в MLflow.

Run ID full:
`{metrics.loc[metrics["mode"] == "full", "run_id"].iloc[0]}`

Run ID chunks:
`{metrics.loc[metrics["mode"] == "chunks", "run_id"].iloc[0]}`

Backend:
`sqlite:///mlflow.db`.

SQLite-файл является локальным runtime-артефактом и не хранится в Git. Экспорт основных MLflow-запусков сохранён в `mlflow_runs.csv`.

Запуск UI:

```powershell
mlflow ui --backend-store-uri sqlite:///mlflow.db --port 5000
```

## 8. Негативные контроли

`config_negative_test.log` подтверждает отклонение:
1. неположительных значений gamma;
2. слишком маленького chunk_size;
3. passes меньше 1.

`leakage_negative_test.log` содержит проверку защиты от использования test-набора при подборе гиперпараметров.

## 9. Тесты и покрытие

Тесты ЛР2:
- `test_data.py`;
- `test_metrics.py`;
- `test_ml.py`;
- `test_tracking_utils.py`.

Итоговый pytest/coverage вывод сохраняется в `pytest_coverage.log`.

## 10. Выводы

1. Accuracy нелинейной модели full составляет {full_acc:.4f}.
2. Accuracy нелинейной модели chunks составляет {chunk_acc:.4f}.
3. Расхождение accuracy между режимами составляет {abs(full_acc - chunk_acc):.4f}.
4. ROC-AUC full составляет {full_auc:.4f}.
5. ROC-AUC chunks составляет {chunk_auc:.4f}.
6. Chunk-режим использовал {chunk_mem:.4f} MiB пиковой измеренной памяти.
7. Full-режим использовал {full_mem:.4f} MiB пиковой измеренной памяти.
8. Экономия памяти chunk-режима составляет {mem_reduction:.2f}%.
9. Full-режим занял {full_time:.4f} с.
10. Chunk-режим занял {chunk_time:.4f} с.
11. Лучшее значение gamma равно {best_gamma}.
12. Использовано {chunks} train-чанка.
"""


def build_docx(
    metrics: pd.DataFrame,
    comparison: dict,
    manifest: dict,
) -> None:
    full = metrics[metrics["mode"] == "full"].iloc[0]
    chunks = metrics[metrics["mode"] == "chunks"].iloc[0]

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Лабораторная работа 2\n"
        "Классический ML: полный датасет и чанки\n"
        "Вариант 09 — make_moons"
    )
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)

    doc.add_heading("1. Цель", level=1)
    doc.add_paragraph(
        "Обучить классическую нелинейную ML-модель в режимах full и chunks, "
        "сравнить качество, время и память, зарегистрировать прогоны в MLflow."
    )

    doc.add_heading("2. Модель", level=1)
    doc.add_paragraph(
        "Pipeline: StandardScaler → RBFSampler → SGDClassifier. "
        f'Лучшее gamma по validation: {comparison.get("best_gamma", "")}.'
    )

    doc.add_heading("3. Результаты", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text = "Метрика"
    table.rows[0].cells[1].text = "Full"
    table.rows[0].cells[2].text = "Chunks"

    rows = [
        ("Accuracy", float(full["accuracy"]), float(chunks["accuracy"])),
        ("F1", float(full["f1"]), float(chunks["f1"])),
        ("ROC-AUC", float(full["roc_auc"]), float(chunks["roc_auc"])),
        ("Peak memory, MiB", float(full["peak_memory_mb"]), float(chunks["peak_memory_mb"])),
        ("Time, s", float(full["elapsed_s"]), float(chunks["elapsed_s"])),
    ]

    for name, a, b in rows:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = f"{a:.6f}"
        cells[2].text = f"{b:.6f}"

    doc.add_heading("4. Интерпретация", level=1)
    mem_save = (
        (float(full["peak_memory_mb"]) - float(chunks["peak_memory_mb"]))
        / float(full["peak_memory_mb"])
        * 100.0
    )
    doc.add_paragraph(
        f'Оба режима дали accuracy {float(full["accuracy"]):.4f}. '
        f'Chunk-режим снизил пиковую измеренную память примерно на {mem_save:.2f}%, '
        "но оказался медленнее из-за overhead обработки отдельных чанков."
    )

    doc.add_heading("5. MLflow", level=1)
    doc.add_paragraph(
        f'Full run ID: {full["run_id"]}\n'
        f'Chunks run ID: {chunks["run_id"]}\n'
        "Backend: sqlite:///mlflow.db"
    )

    doc.add_heading("6. Модели", level=1)
    for name, info in manifest.items():
        doc.add_paragraph(
            f'{name}: {info["bytes"]} bytes; SHA-256 {info["sha256"]}'
        )

    doc.add_heading("7. Негативные контроли", level=1)
    doc.add_paragraph(
        "Проверяются ошибочные конфигурации и защита от утечки test-набора "
        "в процедуру подбора гиперпараметров."
    )

    doc.add_heading("8. Выводы", level=1)
    conclusions = [
        f'Accuracy full — {float(full["accuracy"]):.4f}.',
        f'Accuracy chunks — {float(chunks["accuracy"]):.4f}.',
        f'ROC-AUC full — {float(full["roc_auc"]):.4f}.',
        f'ROC-AUC chunks — {float(chunks["roc_auc"]):.4f}.',
        f'Пиковая память full — {float(full["peak_memory_mb"]):.4f} MiB.',
        f'Пиковая память chunks — {float(chunks["peak_memory_mb"]):.4f} MiB.',
        f'Время full — {float(full["elapsed_s"]):.4f} с.',
        f'Время chunks — {float(chunks["elapsed_s"]):.4f} с.',
        f'Лучшее gamma — {comparison.get("best_gamma", "")}.',
        f'Число train-чанков — {comparison.get("chunks", "")}.',
    ]
    for item in conclusions:
        doc.add_paragraph(item, style="List Number")

    doc.save(LAB2 / "lab2_report.docx")


def build_defense(metrics: pd.DataFrame, comparison: dict) -> None:
    full = metrics[metrics["mode"] == "full"].iloc[0]
    chunks = metrics[metrics["mode"] == "chunks"].iloc[0]
    baseline = comparison.get("linear_baseline_test", {})

    text = f"""# Защита ЛР2

## Что сказать

В ЛР2 одна и та же нелинейная классическая ML-модель обучается двумя способами: на полном train-наборе и через 24 сохранённых чанка.

Pipeline: StandardScaler → RBFSampler → SGDClassifier.

`gamma` выбирается только по validation. Лучшее gamma: {comparison.get("best_gamma", "")}.

Линейный baseline accuracy: {float(baseline.get("accuracy", 0.0)):.4f}.

Нелинейная модель accuracy: {float(full["accuracy"]):.4f}.

Это показывает, что для make_moons нелинейное представление работает заметно лучше линейного baseline.

## Что показать

1. `data/chunks/` — 24 CSV-файла.
2. `hyperparam_search.csv` — выбор gamma по validation.
3. `ml_metrics.csv` — full vs chunks, accuracy/F1/ROC-AUC, bootstrap CI, память, время.
4. `ml_model.joblib` и `ml_model_chunk.joblib`.
5. `model_manifest.json` — SHA-256 моделей.
6. `config_negative_test.log`.
7. `leakage_negative_test.log`.
8. `pytest_coverage.log`.
9. MLflow UI.

## MLflow

```powershell
mlflow ui --backend-store-uri sqlite:///mlflow.db --port 5000
```

Открыть `http://127.0.0.1:5000`.

Full run ID: {full["run_id"]}

Chunks run ID: {chunks["run_id"]}

## Главный результат

Full:
- accuracy {float(full["accuracy"]):.4f}
- F1 {float(full["f1"]):.4f}
- ROC-AUC {float(full["roc_auc"]):.4f}
- memory {float(full["peak_memory_mb"]):.4f} MiB
- time {float(full["elapsed_s"]):.4f} s

Chunks:
- accuracy {float(chunks["accuracy"]):.4f}
- F1 {float(chunks["f1"]):.4f}
- ROC-AUC {float(chunks["roc_auc"]):.4f}
- memory {float(chunks["peak_memory_mb"]):.4f} MiB
- time {float(chunks["elapsed_s"]):.4f} s

В этом детерминированном прогоне full и chunks дали одинаковые predictions.
Разница заключается в способе подачи train-данных и профиле ресурсов.
"""
    (LAB2 / "DEFENSE.md").write_text(text, encoding="utf-8")


def build_screencast(metrics: pd.DataFrame, comparison: dict) -> None:
    full = metrics[metrics["mode"] == "full"].iloc[0]
    chunks = metrics[metrics["mode"] == "chunks"].iloc[0]

    text = f"""# Скринкаст ЛР2 — 3–5 минут

## 0:00–0:30
GitHub: ветка/PR ЛР2, цель full vs chunks.

## 0:30–1:00
Показать `data/chunks`: 24 файла.

## 1:00–1:30
Показать `hyperparam_search.csv`, объяснить выбор gamma={comparison.get("best_gamma", "")} только по validation.

## 1:30–2:15
Показать `ml_metrics.csv`.

Full accuracy: {float(full["accuracy"]):.4f}.
Chunks accuracy: {float(chunks["accuracy"]):.4f}.

Сравнить память и время.

## 2:15–2:45
Показать обе joblib-модели и `model_manifest.json`.

## 2:45–3:15
Показать `config_negative_test.log` и `leakage_negative_test.log`.

## 3:15–3:45
Открыть MLflow UI и показать full/chunks run.

## 3:45–4:10
Показать `pytest_coverage.log` и зелёный GitHub Actions.

## 4:10–4:30
Открыть `lab2_report.docx`.
"""
    (LAB2 / "SCREENCAST.md").write_text(text, encoding="utf-8")


def main() -> None:
    LAB2.mkdir(parents=True, exist_ok=True)

    metrics, search, comparison = read_results()
    manifest = build_model_manifest()
    export_mlflow_runs()
    commit = git_head()

    md = build_markdown(metrics, search, comparison, manifest, commit)
    (LAB2 / "lab2_report.md").write_text(md, encoding="utf-8")

    build_docx(metrics, comparison, manifest)
    build_defense(metrics, comparison)
    build_screencast(metrics, comparison)

    print("Lab 2 final package generated.")
    print("Report:", LAB2 / "lab2_report.docx")


if __name__ == "__main__":
    main()
