from __future__ import annotations

import json
import subprocess
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports" / "LAB1"
FIG_DIR = REPORT_DIR / "figures"
RAW_PATH = ROOT / "data" / "raw" / "moons.csv"


def git_head() -> str:
    return subprocess.check_output(
        ["git", "rev-parse", "HEAD"],
        cwd=ROOT,
        text=True,
    ).strip()


def load_manifest() -> dict:
    return json.loads(
        (REPORT_DIR / "hash_manifest.json").read_text(encoding="utf-8")
    )


def build_schema(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for column in ["x1", "x2", "y"]:
        series = df[column]
        numeric = pd.to_numeric(series, errors="coerce").to_numpy(dtype=float)
        rows.append(
            {
                "field": column,
                "dtype": str(series.dtype),
                "missing": int(series.isna().sum()),
                "non_finite": int((~np.isfinite(numeric)).sum()),
                "unique_values": int(series.nunique()),
                "min": float(series.min()),
                "max": float(series.max()),
            }
        )
    result = pd.DataFrame(rows)
    result.to_csv(REPORT_DIR / "schema.csv", index=False)
    return result


def load_and_extend_metrics(df: pd.DataFrame) -> dict[str, float]:
    path = REPORT_DIR / "eda_stats.csv"
    stats = pd.read_csv(path)
    metrics = {
        str(row["metric"]): float(row["value"])
        for _, row in stats.iterrows()
    }

    numeric = df[["x1", "x2"]].to_numpy(dtype=float)
    metrics["non_finite_feature_values"] = float((~np.isfinite(numeric)).sum())
    metrics["invalid_target_values"] = float((~df["y"].isin([0, 1])).sum())
    metrics["target_unique_values"] = float(df["y"].nunique())

    pd.DataFrame(
        [{"metric": key, "value": value} for key, value in metrics.items()]
    ).to_csv(path, index=False)

    return metrics


def build_markdown(
    df: pd.DataFrame,
    manifest: dict,
    schema: pd.DataFrame,
    metrics: dict[str, float],
    commit: str,
) -> str:
    class_counts = df["y"].value_counts().sort_index()
    x1_dtype = schema.loc[schema["field"] == "x1", "dtype"].iloc[0]
    x2_dtype = schema.loc[schema["field"] == "x2", "dtype"].iloc[0]
    y_dtype = schema.loc[schema["field"] == "y", "dtype"].iloc[0]

    return f"""# Лабораторная работа 1. Git-окружение, pre-commit и EDA

## 1. Паспорт данных

Автор: Vladimir

Вариант: 09

Тема варианта: `make_moons` — нелинейность.

Хеш коммита кода: `{commit}`.

Источник: `sklearn.datasets.make_moons`.

Для варианта 09 внешний CSV не скачивается: данные воспроизводимо генерируются встроенным генератором scikit-learn. Библиотека scikit-learn распространяется по BSD-3-Clause; отдельной лицензии внешнего датасета нет, поскольку исходный CSV создаётся локально.

Параметры генерации:

- `n_samples = {manifest["generator"]["n_samples"]}`;
- `noise = {manifest["generator"]["noise"]}`;
- `seed = {manifest["generator"]["seed"]}`.

Raw-файл: `{manifest["raw"]["path"]}`.

Размер raw-файла: {manifest["raw"]["bytes"]} байт.

SHA-256:

`{manifest["raw"]["sha256"]}`

## 2. Объём и структура

Объём: {len(df)} строк.

| Поле | Тип | Назначение |
|---|---|---|
| x1 | {x1_dtype} | первый числовой признак |
| x2 | {x2_dtype} | второй числовой признак |
| y | {y_dtype} | бинарный целевой класс |

Пропущенных ячеек: {int(df.isna().sum().sum())}.

Дублирующихся строк: {int(df.duplicated().sum())}.

Нечисловых или бесконечных значений признаков: {int(metrics["non_finite_feature_values"])}.

Недопустимых значений y: {int(metrics["invalid_target_values"])}.

Минимумы и максимумы полей сохранены в `schema.csv`. Экстремальные, но конечные значения признаков не считаются ошибкой автоматически, поскольку данные генерируются с шумом.

## 3. План разбиения

Train: {manifest["splits"]["train"]}.

Validation: {manifest["splits"]["validation"]}.

Test: {manifest["splits"]["test"]}.

Размер чанка: {manifest["generator"]["chunk_size"]}.

Train-чанков: {manifest["splits"]["chunks"]}.

## 4. EDA

Машиночитаемые сводки: `eda_stats.csv`.

Класс 0: {int(class_counts.loc[0])} объектов.

Класс 1: {int(class_counts.loc[1])} объектов.

Доля класса 0: 50.0 %.

Доля класса 1: 50.0 %.

Среднее x1: {metrics["x1_mean"]:.6f}.

Стандартное отклонение x1: {metrics["x1_std"]:.6f}.

Среднее x2: {metrics["x2_mean"]:.6f}.

Стандартное отклонение x2: {metrics["x2_std"]:.6f}.

Корреляция x1 и x2: {metrics["x1_x2_corr"]:.6f}.

## 5. Графики

Главный график: `figures/nonlinearity_scatter.png`.

Он показывает два полумесяца и линейную границу LogisticRegression. Геометрия классов нелинейная, поэтому одна прямая не может полностью разделить выборку.

Диагностическая accuracy линейной границы на полном raw-наборе: {metrics["linear_train_accuracy"]:.4f}.

Эта accuracy не является итоговой оценкой ML-модели и используется только как EDA-ориентир.

Дополнительные графики:

- `class_balance.png` — баланс классов;
- `feature_distributions.png` — распределения x1 и x2 по классам;
- `correlation_matrix.png` — корреляции x1, x2 и y;
- `linear_confusion.png` — матрица ошибок линейного ориентира.

## 6. Pre-commit

Настроены:

1. проверка синтаксиса и базового стиля Python;
2. поиск секретов;
3. запрет файлов больше 5 MiB.

`precommit_block.log` содержит реальные неудачные попытки `git commit` для плохого Python-файла, тестового секрета и большого бинарного файла.

`precommit_clean.log` содержит успешную чистую проверку репозитория.

## 7. Контроль целостности и воспроизводимость

`hash_manifest.json` содержит путь, размер и SHA-256 raw-файла.

`hash_negative_test.log` показывает, что после подмены файла SHA-256 меняется и подмена обнаруживается.

`reproducibility.log` показывает, что повторная генерация с тем же seed даёт тот же SHA-256.

## 8. Нумерованные выводы

1. Объём исходной выборки составляет {len(df)} объектов.
2. Число входных признаков равно 2.
3. Число пропущенных значений равно {int(df.isna().sum().sum())}.
4. Число дублирующихся строк равно {int(df.duplicated().sum())}.
5. Доля класса 0 составляет 50.0 %.
6. Доля класса 1 составляет 50.0 %.
7. Корреляция x1 и x2 равна {metrics["x1_x2_corr"]:.4f}.
8. Accuracy линейного ориентира равна {metrics["linear_train_accuracy"]:.4f}.
9. Train содержит {manifest["splits"]["train"]} объектов.
10. Validation содержит {manifest["splits"]["validation"]} объектов.
11. Test содержит {manifest["splits"]["test"]} объектов.
12. Подготовлено {manifest["splits"]["chunks"]} train-чанка.
"""


def add_picture_with_caption(
    document: Document,
    path: Path,
    caption: str,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(6.0))

    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)


def build_docx(
    manifest: dict,
    metrics: dict[str, float],
    commit: str,
) -> None:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "Лабораторная работа 1\n"
        "Git-окружение, pre-commit и EDA\n"
        "Вариант 09 — make_moons: нелинейность"
    )
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    document.add_heading("1. Паспорт данных", level=1)

    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    passport = [
        ("Автор", "Vladimir"),
        ("Вариант", "09"),
        ("Источник", "sklearn.datasets.make_moons"),
        ("Происхождение", "локальная генерация scikit-learn"),
        ("Лицензия библиотеки", "BSD-3-Clause"),
        ("Объём", "24000"),
        ("SHA-256", manifest["raw"]["sha256"]),
        ("Хеш коммита", commit),
    ]

    for key, value in passport:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value

    document.add_heading("2. Структура и качество данных", level=1)
    document.add_paragraph(
        "Поля: x1 (float64), x2 (float64), y (int64). "
        "Пропущенных значений: 0. Дубликатов: 0. "
        "Нечисловых или бесконечных значений признаков: 0. "
        "Недопустимых значений целевого класса: 0."
    )

    document.add_heading("3. Разбиение", level=1)
    document.add_paragraph(
        f'Train: {manifest["splits"]["train"]}; '
        f'validation: {manifest["splits"]["validation"]}; '
        f'test: {manifest["splits"]["test"]}; '
        f'chunks: {manifest["splits"]["chunks"]}; '
        f'chunk size: {manifest["generator"]["chunk_size"]}.'
    )

    document.add_heading("4. EDA", level=1)
    eda = document.add_table(rows=1, cols=2)
    eda.style = "Table Grid"
    eda.rows[0].cells[0].text = "Показатель"
    eda.rows[0].cells[1].text = "Значение"

    values = [
        ("Объекты", "24000"),
        ("Класс 0", "12000 (50.0 %)"),
        ("Класс 1", "12000 (50.0 %)"),
        ("corr(x1, x2)", f'{metrics["x1_x2_corr"]:.4f}'),
        ("Linear accuracy", f'{metrics["linear_train_accuracy"]:.4f}'),
    ]

    for key, value in values:
        cells = eda.add_row().cells
        cells[0].text = key
        cells[1].text = value

    document.add_heading("5. Визуализации", level=1)

    figures = [
        ("nonlinearity_scatter.png",
         "Рисунок 1 — нелинейность классов и линейная граница"),
        ("class_balance.png",
         "Рисунок 2 — баланс классов"),
        ("feature_distributions.png",
         "Рисунок 3 — распределения признаков"),
        ("correlation_matrix.png",
         "Рисунок 4 — корреляционная матрица"),
        ("linear_confusion.png",
         "Рисунок 5 — матрица ошибок линейного ориентира"),
    ]

    for filename, caption in figures:
        path = FIG_DIR / filename
        if path.exists():
            add_picture_with_caption(document, path, caption)

    document.add_heading("6. Интерпретация", level=1)
    document.add_paragraph(
        "Классы образуют два полумесяца. Одна линейная граница "
        "не соответствует этой геометрии. Диагностическая "
        f'LogisticRegression имеет accuracy {metrics["linear_train_accuracy"]:.4f}. '
        "Эта величина используется только как EDA-ориентир."
    )

    document.add_heading("7. Негативные контроли", level=1)
    document.add_paragraph(
        "Реальные попытки git commit с синтаксически неверным Python-файлом, "
        "тестовым секретом и бинарным файлом больше 5 MiB блокируются "
        "pre-commit. Подмена raw-файла обнаруживается по изменению SHA-256."
    )

    document.add_heading("8. Воспроизводимость", level=1)
    document.add_paragraph(
        "Повторная генерация с теми же параметрами и seed даёт тот же SHA-256."
    )

    document.add_heading("9. Выводы", level=1)
    conclusions = [
        "Объём выборки — 24000 объектов.",
        "Число входных признаков — 2.",
        "Число пропущенных значений — 0.",
        "Число дубликатов — 0.",
        "Доля класса 0 — 50.0 %.",
        "Доля класса 1 — 50.0 %.",
        f'Корреляция x1 и x2 — {metrics["x1_x2_corr"]:.4f}.',
        f'Accuracy линейного ориентира — {metrics["linear_train_accuracy"]:.4f}.',
        f'Количество train-чанков — {manifest["splits"]["chunks"]}.',
    ]
    for item in conclusions:
        document.add_paragraph(item, style="List Number")

    document.save(REPORT_DIR / "lab1_report.docx")


def build_defense(
    manifest: dict,
    metrics: dict[str, float],
) -> None:
    text = rf"""# Защита ЛР1

## Короткое вступление

У меня вариант 09 — make_moons. Основная проблема — нелинейность.

Внешний CSV для этого варианта не скачивается: make_moons является встроенным генератором scikit-learn. Я фиксирую параметры генерации, сохраняю raw CSV и его SHA-256.

## Основные числа

- объектов: 24000;
- train: {manifest["splits"]["train"]};
- validation: {manifest["splits"]["validation"]};
- test: {manifest["splits"]["test"]};
- чанков: {manifest["splits"]["chunks"]};
- классы: 50 % / 50 %;
- corr(x1,x2): {metrics["x1_x2_corr"]:.4f};
- accuracy линейного ориентира: {metrics["linear_train_accuracy"]:.4f}.

## Порядок демонстрации

1. GitHub: репозиторий, последние коммиты, Actions, Protect main.
2. `hash_manifest.json`: источник, параметры, размер, SHA-256.
3. `eda_stats.csv` и `schema.csv`.
4. `nonlinearity_scatter.png`: два полумесяца и линейная граница.
5. Остальные графики: баланс, распределения, корреляции, confusion matrix.
6. `precommit_block.log`: три заблокированных git commit.
7. `hash_negative_test.log`: подмена данных обнаружена.
8. `reproducibility.log`: повторная генерация дала тот же SHA-256.
9. `tests.log`: 3 passed.
10. `lab1_report.docx`.

## Что говорить про главный график

Два класса образуют два полумесяца. Одна прямая граница не может полностью разделить такую геометрию. Поэтому заявленная проблема варианта — нелинейность — видна непосредственно на данных.

Accuracy 0.8679 — это не итоговая оценка ML-модели, а диагностический линейный ориентир EDA.

## Команды на защите

```powershell
git log --oneline -5
git status
Get-Content .\reports\LAB1\hash_manifest.json
Get-Content .\reports\LAB1\eda_stats.csv
Get-Content .\reports\LAB1\precommit_block.log
Get-Content .\reports\LAB1\hash_negative_test.log
Get-Content .\reports\LAB1\reproducibility.log
python -m pytest .\tests\test_data.py .\tests\test_eda.py -q
explorer .\reports\LAB1\figures
```
"""
    (REPORT_DIR / "DEFENSE.md").write_text(text, encoding="utf-8")


def build_screencast(
    manifest: dict,
    metrics: dict[str, float],
) -> None:
    text = rf"""# Скринкаст ЛР1 — план на 3–5 минут

## 0:00–0:30
Открыть GitHub. Назвать вариант 09, make_moons и проблему нелинейности.

## 0:30–1:00
Показать `hash_manifest.json`: 24000 объектов, seed 42, noise 0.22, SHA-256.

## 1:00–1:40
Показать `eda_stats.csv` и `schema.csv`: пропусков 0, дубликатов 0, классы 50/50.

## 1:40–2:30
Показать `nonlinearity_scatter.png`.
Объяснить два полумесяца и линейную границу.
Linear accuracy: {metrics["linear_train_accuracy"]:.4f}.

Быстро показать остальные четыре графика.

## 2:30–3:10
Показать `precommit_block.log`: три реальные попытки commit заблокированы.

## 3:10–3:35
Показать `hash_negative_test.log` и `reproducibility.log`.

## 3:35–4:00
Запустить:
`python -m pytest .\\tests\\test_data.py .\\tests\\test_eda.py -q`

Показать `3 passed`.

## 4:00–4:20
Показать зелёные GitHub Actions и активный `Protect main`.

## 4:20–4:35
Открыть `lab1_report.docx`.
"""
    (REPORT_DIR / "SCREENCAST.md").write_text(text, encoding="utf-8")


def main() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)

    df = pd.read_csv(RAW_PATH)
    manifest = load_manifest()
    schema = build_schema(df)
    metrics = load_and_extend_metrics(df)
    commit = git_head()

    markdown = build_markdown(df, manifest, schema, metrics, commit)
    (REPORT_DIR / "lab1_report.md").write_text(markdown, encoding="utf-8")

    build_docx(manifest, metrics, commit)
    build_defense(manifest, metrics)
    build_screencast(manifest, metrics)

    print("Final Lab 1 package generated.")
    print("DOCX:", REPORT_DIR / "lab1_report.docx")
    print("Defense:", REPORT_DIR / "DEFENSE.md")
    print("Screencast:", REPORT_DIR / "SCREENCAST.md")


if __name__ == "__main__":
    main()

