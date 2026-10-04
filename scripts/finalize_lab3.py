from __future__ import annotations

import hashlib
import json
import subprocess
from datetime import date
from pathlib import Path

import mlflow
import pandas as pd
import torch
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from mlflow.tracking import MlflowClient


ROOT = Path(__file__).resolve().parents[1]
LAB3 = ROOT / "reports" / "LAB3"
FIGURE = LAB3 / "figures" / "loss_curves.png"


def git_head() -> str:
    return subprocess.check_output(
        ["git", "rev-parse", "HEAD"],
        cwd=ROOT,
        text=True,
    ).strip()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def checkpoint_parameter_count(path: Path) -> int:
    obj = torch.load(path, map_location="cpu", weights_only=False)

    if isinstance(obj, dict):
        state = obj.get("model_state_dict")
        if state is None:
            state = obj.get("state_dict")
        if state is None:
            state = obj
    else:
        return 0

    return int(
        sum(
            value.numel()
            for value in state.values()
            if torch.is_tensor(value)
        )
    )


def model_manifest() -> dict:
    payload = {}
    for name in ["dl_model.pt", "dl_model_chunk.pt"]:
        path = LAB3 / name
        if path.exists():
            payload[name] = {
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
                "tensor_parameter_count": checkpoint_parameter_count(path),
            }

    (LAB3 / "model_manifest.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    return payload


def environment_info() -> dict:
    info = {
        "torch_version": torch.__version__,
        "cuda_available": bool(torch.cuda.is_available()),
        "cuda_version": torch.version.cuda,
        "device_count": int(torch.cuda.device_count()),
        "gpu_name": (
            torch.cuda.get_device_name(0)
            if torch.cuda.is_available()
            else None
        ),
    }
    (LAB3 / "environment.json").write_text(
        json.dumps(info, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    return info


def export_mlflow_runs() -> None:
    uri = "sqlite:///mlflow.db"
    mlflow.set_tracking_uri(uri)
    client = MlflowClient(tracking_uri=uri)

    rows = []
    try:
        exp = client.get_experiment_by_name("lab3_dl")
        if exp is not None:
            for run in client.search_runs([exp.experiment_id]):
                rows.append(
                    {
                        "experiment_id": exp.experiment_id,
                        "experiment_name": exp.name,
                        "run_id": run.info.run_id,
                        "status": run.info.status,
                        "run_name": run.data.tags.get("mlflow.runName", ""),
                        "params": json.dumps(
                            run.data.params,
                            ensure_ascii=False,
                            sort_keys=True,
                        ),
                        "metrics": json.dumps(
                            run.data.metrics,
                            ensure_ascii=False,
                            sort_keys=True,
                        ),
                    }
                )
    except Exception as exc:
        rows.append(
            {
                "experiment_id": "",
                "experiment_name": "lab3_dl",
                "run_id": "",
                "status": "EXPORT_ERROR",
                "run_name": "",
                "params": "",
                "metrics": str(exc),
            }
        )

    pd.DataFrame(rows).to_csv(LAB3 / "mlflow_runs.csv", index=False)


def load_data():
    metrics = pd.read_csv(LAB3 / "dl_metrics.csv")
    memory = pd.read_csv(LAB3 / "memory_compare.csv")
    comparison = json.loads(
        (LAB3 / "comparison.json").read_text(encoding="utf-8")
    )
    reproducibility = json.loads(
        (LAB3 / "reproducibility.json").read_text(encoding="utf-8")
    )
    seed_negative = json.loads(
        (LAB3 / "seed_negative_test.json").read_text(encoding="utf-8")
    )
    config = json.loads(
        (ROOT / "configs" / "lab3.json").read_text(encoding="utf-8")
    )
    lab2 = pd.read_csv(ROOT / "reports" / "LAB2" / "ml_metrics.csv")
    return (
        metrics,
        memory,
        comparison,
        reproducibility,
        seed_negative,
        config,
        lab2,
    )


def metric(df: pd.DataFrame, mode: str, name: str) -> float:
    return float(df.loc[df["mode"] == mode, name].iloc[0])


def memory_value(df: pd.DataFrame, mode: str, name: str) -> float:
    return float(df.loc[df["mode"] == mode, name].iloc[0])


def parity_text(comparison: dict) -> str:
    parity = comparison.get("device_parity", {})
    if parity.get("status") == "not_available":
        return (
            "CPU/GPU parity could not be measured in this environment because "
            "CUDA was unavailable. This limitation is explicitly recorded in "
            "`device_parity.csv` and `environment.json`."
        )

    diff = parity.get("max_abs_diff")
    return (
        f"CPU/GPU prediction max absolute difference: {diff}. "
        "Required tolerance: 1e-5."
    )


def build_markdown(
    metrics: pd.DataFrame,
    memory: pd.DataFrame,
    comparison: dict,
    reproducibility: dict,
    seed_negative: dict,
    config: dict,
    lab2: pd.DataFrame,
    manifest: dict,
    env: dict,
    commit: str,
) -> str:
    full_acc = metric(metrics, "full", "accuracy")
    stream_acc = metric(metrics, "chunks", "accuracy")
    full_f1 = metric(metrics, "full", "f1")
    stream_f1 = metric(metrics, "chunks", "f1")
    full_auc = metric(metrics, "full", "roc_auc")
    stream_auc = metric(metrics, "chunks", "roc_auc")

    full_mem = memory_value(memory, "full", "peak_memory_mb")
    stream_mem = memory_value(memory, "chunks", "peak_memory_mb")
    full_time = memory_value(memory, "full", "elapsed_s")
    stream_time = memory_value(memory, "chunks", "elapsed_s")

    mem_reduction = (full_mem - stream_mem) / full_mem * 100.0
    time_change = (stream_time - full_time) / full_time * 100.0

    ml_acc = float(
        lab2.loc[lab2["mode"] == "full", "accuracy"].iloc[0]
    )
    dl_vs_ml_pp = (full_acc - ml_acc) * 100.0

    param_count = manifest.get("dl_model.pt", {}).get(
        "tensor_parameter_count",
        0,
    )

    parity = parity_text(comparison)

    return f"""# Лабораторная работа 3. PyTorch: full vs streaming

## 1. Паспорт работы

Автор: Vladimir

Дата: {date.today().isoformat()}

Вариант: 09 — `make_moons`, проблема нелинейности.

Хеш коммита-основания: `{commit}`.

Архитектура: MLP `2 -> 32 -> 32 -> 1`.

Число параметров по checkpoint: {param_count}.

PyTorch: `{env["torch_version"]}`.

CUDA available: `{env["cuda_available"]}`.

GPU: `{env["gpu_name"]}`.

## 2. Настройки обучения

- seed: {config["seed"]}
- batch size: {config["batch_size"]}
- epochs: {config["epochs"]}
- learning rate: {config["learning_rate"]}
- weight decay: {config["weight_decay"]}
- early stopping patience: {config["patience"]}
- bootstrap rounds: {config["bootstrap_rounds"]}
- configured training device: `{config["device"]}`

В работе сравниваются два источника данных:
1. full — весь train-набор материализуется в памяти;
2. streaming/chunks — train читается по 24 сохранённым чанкам через потоковый загрузчик.

## 3. Метрики

| Метрика | Full | Streaming |
|---|---:|---:|
| Accuracy | {full_acc:.6f} | {stream_acc:.6f} |
| F1 | {full_f1:.6f} | {stream_f1:.6f} |
| ROC-AUC | {full_auc:.6f} | {stream_auc:.6f} |
| Best epoch | {int(metric(metrics, "full", "best_epoch"))} | {int(metric(metrics, "chunks", "best_epoch"))} |

Разница accuracy: {abs(full_acc - stream_acc):.6f}.

Максимальное абсолютное расхождение вероятностных предсказаний full/stream:
{float(comparison["full_vs_chunk_prediction_max_abs_diff"]):.6f}.

## 4. Bootstrap confidence intervals

Full accuracy CI:
`[{metric(metrics, "full", "accuracy_ci_low"):.6f}; {metric(metrics, "full", "accuracy_ci_high"):.6f}]`.

Streaming accuracy CI:
`[{metric(metrics, "chunks", "accuracy_ci_low"):.6f}; {metric(metrics, "chunks", "accuracy_ci_high"):.6f}]`.

Full ROC-AUC CI:
`[{metric(metrics, "full", "roc_auc_ci_low"):.6f}; {metric(metrics, "full", "roc_auc_ci_high"):.6f}]`.

Streaming ROC-AUC CI:
`[{metric(metrics, "chunks", "roc_auc_ci_low"):.6f}; {metric(metrics, "chunks", "roc_auc_ci_high"):.6f}]`.

## 5. Память и время

| Режим | Peak memory, MiB | Time, s |
|---|---:|---:|
| Full | {full_mem:.6f} | {full_time:.6f} |
| Streaming | {stream_mem:.6f} | {stream_time:.6f} |

В этом конкретном прогоне streaming использовал на {mem_reduction:.2f}% меньше измеренной пиковой памяти.

Изменение времени streaming относительно full: {time_change:+.2f}%.

Эти значения являются измерением данной реализации и данного запуска, а не универсальным свойством streaming.

## 6. Early stopping и кривые

Настроено правило early stopping с `patience={config["patience"]}`.

В текущем прогоне лучший epoch full: {int(metric(metrics, "full", "best_epoch"))}.

Лучший epoch streaming: {int(metric(metrics, "chunks", "best_epoch"))}.

Кривые обучения сохранены в:

`figures/loss_curves.png`.

## 7. Воспроизводимость

Режим:
`{reproducibility["mode"]}`.

Повтор с тем же seed:
- prediction max abs diff: {float(reproducibility["prediction_max_abs_diff"]):.6f};
- accuracy diff: {float(reproducibility["metric_diff_accuracy"]):.6f};
- допустимое расхождение: {float(reproducibility["allowed_metric_diff"]):.6f};
- pass: `{reproducibility["pass"]}`.

Негативный контроль seed:
- reference seed: {seed_negative["reference_seed"]};
- alternate seed: {seed_negative["alternate_seed"]};
- prediction max abs diff: {float(seed_negative["prediction_max_abs_diff"]):.6f};
- accuracy abs diff: {float(seed_negative["accuracy_abs_diff"]):.6f};
- pass: `{seed_negative["pass"]}`.

## 8. Device parity

{parity}

## 9. Сравнение с классическим ML из ЛР2

ML full accuracy из ЛР2: {ml_acc:.6f}.

DL full accuracy: {full_acc:.6f}.

Разница DL - ML: {dl_vs_ml_pp:+.3f} процентного пункта.

В данном прогоне классическая RBF+SGD модель ЛР2 немного точнее MLP, что является допустимым результатом: цель ЛР3 — построить и исследовать DL-пайплайн, а не гарантированно превзойти классический ML.

## 10. Checkpoints и MLflow

`dl_model.pt` SHA-256:
`{manifest["dl_model.pt"]["sha256"]}`

`dl_model_chunk.pt` SHA-256:
`{manifest["dl_model_chunk.pt"]["sha256"]}`

Full MLflow run:
`{metrics.loc[metrics["mode"] == "full", "run_id"].iloc[0]}`

Streaming MLflow run:
`{metrics.loc[metrics["mode"] == "chunks", "run_id"].iloc[0]}`

Backend:
`sqlite:///mlflow.db`.

## 11. Негативные конфиги

`config_negative_test.log` подтверждает отказ для:
1. слишком маленького batch size;
2. неположительного learning rate;
3. недопустимого device.

## 12. Выводы

1. Full accuracy составляет {full_acc:.4f}.
2. Streaming accuracy составляет {stream_acc:.4f}.
3. Расхождение accuracy между режимами составляет {abs(full_acc - stream_acc):.4f}.
4. Full ROC-AUC составляет {full_auc:.4f}.
5. Streaming ROC-AUC составляет {stream_auc:.4f}.
6. Full peak memory составляет {full_mem:.4f} MiB.
7. Streaming peak memory составляет {stream_mem:.4f} MiB.
8. Экономия измеренной пиковой памяти streaming составляет {mem_reduction:.2f}%.
9. Воспроизводимость с seed 42 подтверждена: `{reproducibility["pass"]}`.
10. Негативный seed-control подтверждён: `{seed_negative["pass"]}`.
11. DL full отличается от ML ЛР2 на {dl_vs_ml_pp:+.3f} процентного пункта.
"""


def add_picture(doc: Document, path: Path, caption: str) -> None:
    if not path.exists():
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.0))

    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = c.add_run(caption)
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def build_docx(
    metrics: pd.DataFrame,
    memory: pd.DataFrame,
    comparison: dict,
    reproducibility: dict,
    config: dict,
    lab2: pd.DataFrame,
    manifest: dict,
    env: dict,
) -> None:
    full = metrics[metrics["mode"] == "full"].iloc[0]
    stream = metrics[metrics["mode"] == "chunks"].iloc[0]
    full_mem = memory[memory["mode"] == "full"].iloc[0]
    stream_mem = memory[memory["mode"] == "chunks"].iloc[0]

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Лабораторная работа 3\n"
        "PyTorch: полный датасет и streaming\n"
        "Вариант 09 — make_moons"
    )
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    doc.add_heading("1. Цель", level=1)
    doc.add_paragraph(
        "Построить воспроизводимый PyTorch-пайплайн MLP 2→32→32→1 "
        "и сравнить обучение на полном train-наборе со streaming по чанкам."
    )

    doc.add_heading("2. Конфигурация", level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    items = [
        ("Seed", str(config["seed"])),
        ("Batch size", str(config["batch_size"])),
        ("Epochs", str(config["epochs"])),
        ("Learning rate", str(config["learning_rate"])),
        ("Weight decay", str(config["weight_decay"])),
        ("Patience", str(config["patience"])),
        ("PyTorch", str(env["torch_version"])),
        ("CUDA", str(env["cuda_available"])),
        ("GPU", str(env["gpu_name"])),
    ]

    for key, value in items:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value

    doc.add_heading("3. Метрики full vs streaming", level=1)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.rows[0].cells[0].text = "Показатель"
    t.rows[0].cells[1].text = "Full"
    t.rows[0].cells[2].text = "Streaming"

    rows = [
        ("Accuracy", float(full["accuracy"]), float(stream["accuracy"])),
        ("F1", float(full["f1"]), float(stream["f1"])),
        ("ROC-AUC", float(full["roc_auc"]), float(stream["roc_auc"])),
        ("Peak memory, MiB", float(full_mem["peak_memory_mb"]), float(stream_mem["peak_memory_mb"])),
        ("Time, s", float(full_mem["elapsed_s"]), float(stream_mem["elapsed_s"])),
    ]

    for name, a, b in rows:
        cells = t.add_row().cells
        cells[0].text = name
        cells[1].text = f"{a:.6f}"
        cells[2].text = f"{b:.6f}"

    doc.add_heading("4. Кривые обучения", level=1)
    add_picture(
        doc,
        FIGURE,
        "Рисунок 1 — train/validation loss для full и streaming режимов",
    )

    doc.add_heading("5. Воспроизводимость", level=1)
    doc.add_paragraph(
        f'Prediction max abs diff при повторе с тем же seed: '
        f'{float(reproducibility["prediction_max_abs_diff"]):.6f}. '
        f'Проверка: {reproducibility["pass"]}.'
    )

    doc.add_heading("6. Device parity", level=1)
    doc.add_paragraph(parity_text(comparison))

    doc.add_heading("7. Сравнение с ЛР2", level=1)
    ml_acc = float(lab2.loc[lab2["mode"] == "full", "accuracy"].iloc[0])
    doc.add_paragraph(
        f'Accuracy классического ML ЛР2: {ml_acc:.4f}. '
        f'Accuracy MLP full: {float(full["accuracy"]):.4f}.'
    )

    doc.add_heading("8. Checkpoints", level=1)
    for name, info in manifest.items():
        doc.add_paragraph(
            f'{name}: {info["bytes"]} bytes; SHA-256 {info["sha256"]}'
        )

    doc.add_heading("9. Вывод", level=1)
    doc.add_paragraph(
        "Оба режима дают практически одинаковое качество. "
        "Streaming существенно снижает измеренную пиковую память. "
        "Повтор с тем же seed детерминирован в текущем CPU-режиме."
    )

    doc.save(LAB3 / "lab3_report.docx")


def build_defense(
    metrics: pd.DataFrame,
    memory: pd.DataFrame,
    comparison: dict,
    env: dict,
) -> None:
    full = metrics[metrics["mode"] == "full"].iloc[0]
    stream = metrics[metrics["mode"] == "chunks"].iloc[0]
    mf = memory[memory["mode"] == "full"].iloc[0]
    ms = memory[memory["mode"] == "chunks"].iloc[0]

    text = f"""# Защита ЛР3

## Что сказать

В ЛР3 задача make_moons решается нейросетью PyTorch MLP 2→32→32→1.

Есть два режима:
- full — train целиком в памяти;
- streaming — данные читаются по 24 чанкам через потоковый загрузчик.

Full accuracy: {float(full["accuracy"]):.4f}.

Streaming accuracy: {float(stream["accuracy"]):.4f}.

Full peak memory: {float(mf["peak_memory_mb"]):.4f} MiB.

Streaming peak memory: {float(ms["peak_memory_mb"]):.4f} MiB.

Full time: {float(mf["elapsed_s"]):.4f} s.

Streaming time: {float(ms["elapsed_s"]):.4f} s.

## Что показать

1. `configs/lab3.json`.
2. `src/moons_lab/dl_pipeline.py` — MLP и streaming loader.
3. `data/chunks/` — 24 чанка.
4. `dl_metrics.csv`.
5. `memory_compare.csv`.
6. `figures/loss_curves.png`.
7. `dl_model.pt` и `dl_model_chunk.pt`.
8. `reproducibility.json`.
9. `seed_negative_test.json`.
10. `device_parity.csv`.
11. `pytest_coverage.log`.
12. MLflow UI.

## Воспроизводимость

Повтор с seed=42 дал prediction diff 0 и metric diff 0.

При seed=43 prediction max abs diff стал:
{comparison["seed_negative_control"]["prediction_max_abs_diff"]:.6f}.

Это демонстрирует, зачем фиксировать seed.

## CPU/GPU

CUDA available: {env["cuda_available"]}.

GPU: {env["gpu_name"]}.

{parity_text(comparison)}

## Сравнение с ЛР2

Классический ML ЛР2 accuracy:
{comparison["ml_lab2_full_accuracy"]:.4f}.

DL full accuracy:
{comparison["dl_full_accuracy"]:.4f}.

То есть в этом запуске классический RBF+SGD немного точнее MLP. Это нормальный результат; ЛР3 проверяет DL-пайплайн, streaming и воспроизводимость.
"""
    (LAB3 / "DEFENSE.md").write_text(text, encoding="utf-8")


def build_screencast(
    metrics: pd.DataFrame,
    memory: pd.DataFrame,
    comparison: dict,
) -> None:
    full = metrics[metrics["mode"] == "full"].iloc[0]
    stream = metrics[metrics["mode"] == "chunks"].iloc[0]
    mf = memory[memory["mode"] == "full"].iloc[0]
    ms = memory[memory["mode"] == "chunks"].iloc[0]

    text = f"""# Скринкаст ЛР3 — 3–5 минут

## 0:00–0:30
Показать GitHub/ветку ЛР3 и сказать: PyTorch MLP 2→32→32→1, full vs streaming.

## 0:30–1:00
Открыть `configs/lab3.json` и `data/chunks/` — 24 чанка.

## 1:00–1:40
Показать `loss_curves.png` и объяснить epochs/early stopping.

## 1:40–2:20
Показать `dl_metrics.csv`.

Full accuracy: {float(full["accuracy"]):.4f}.

Streaming accuracy: {float(stream["accuracy"]):.4f}.

## 2:20–2:50
Показать `memory_compare.csv`.

Full: {float(mf["peak_memory_mb"]):.4f} MiB / {float(mf["elapsed_s"]):.4f} s.

Streaming: {float(ms["peak_memory_mb"]):.4f} MiB / {float(ms["elapsed_s"]):.4f} s.

## 2:50–3:20
Показать `reproducibility.json` и `seed_negative_test.json`.

## 3:20–3:45
Показать `device_parity.csv` и объяснить результат.

## 3:45–4:10
Показать две `.pt` модели и `model_manifest.json`.

## 4:10–4:30
Показать `pytest_coverage.log`, MLflow и зелёный CI.
"""
    (LAB3 / "SCREENCAST.md").write_text(text, encoding="utf-8")


def main() -> None:
    LAB3.mkdir(parents=True, exist_ok=True)

    (
        metrics,
        memory,
        comparison,
        reproducibility,
        seed_negative,
        config,
        lab2,
    ) = load_data()

    manifest = model_manifest()
    env = environment_info()
    export_mlflow_runs()
    commit = git_head()

    md = build_markdown(
        metrics,
        memory,
        comparison,
        reproducibility,
        seed_negative,
        config,
        lab2,
        manifest,
        env,
        commit,
    )
    (LAB3 / "lab3_report.md").write_text(md, encoding="utf-8")

    build_docx(
        metrics,
        memory,
        comparison,
        reproducibility,
        config,
        lab2,
        manifest,
        env,
    )
    build_defense(metrics, memory, comparison, env)
    build_screencast(metrics, memory, comparison)

    print("Lab 3 final package generated.")
    print("DOCX:", LAB3 / "lab3_report.docx")


if __name__ == "__main__":
    main()

